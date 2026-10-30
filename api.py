import os
import io
import re
import copy
from datetime import datetime, timedelta
from typing import Optional, Dict, Any, List
from fastapi import FastAPI, HTTPException, Depends, Query, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import HTMLResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from sqlalchemy.orm import Session
from pydantic import BaseModel
from database import init_db, get_db, User, SavedTimetable, Substitution
from auth import get_password_hash, verify_password, create_access_token, get_current_user
from models import TimetableRequest, TimetableResponse, AIInputRequest
from solver import validate_teachers, solve_timetable
from excel_generator import generate_class_template_excel
from pdf_generator import generate_teachers_pdf, generate_classes_pdf


# ============================================================
# PYDANTIC МОДЕЛИ
# ============================================================
class UserRegister(BaseModel):
    email: str
    password: str
    full_name: str
    school_name: Optional[str] = None


class UserLogin(BaseModel):
    email: str
    password: str


class TimetableUpdate(BaseModel):
    name: Optional[str] = None
    data: Optional[Dict[str, Any]] = None


class TopicUpdate(BaseModel):
    teacher: str
    subject: str
    grade: str
    day: str
    slot: int
    topic: str


class TopicsUpdateRequest(BaseModel):
    updates: List[TopicUpdate]


# ============================================================
# ВСПОМОГАТЕЛЬНЫЕ ФУНКЦИИ
# ============================================================
def ensure_topic_field(data: Dict[str, Any]) -> Dict[str, Any]:
    """Гарантирует наличие поля topic во всех уроках"""
    if not data or 'teachers' not in data:
        return data

    for teacher in data['teachers']:
        if 'schedule' not in teacher:
            continue
        for day, day_data in teacher['schedule'].items():
            for slot, lessons in day_data.items():
                if isinstance(lessons, list):
                    for lesson in lessons:
                        if 'topic' not in lesson:
                            lesson['topic'] = ''
    return data


def remove_duplicates_from_schedule(schedule: Dict) -> Dict:
    """Удаляет дубликаты уроков в расписании"""
    for day_key, day_data in schedule.items():
        for slot_key, lessons in day_data.items():
            if isinstance(lessons, list):
                seen = set()
                unique_lessons = []
                for lesson in lessons:
                    # Ключ включает subject, grade, teacher, classroom
                    key = f"{lesson.get('subject', '')}_{lesson.get('grade', '')}_{lesson.get('teacher', '')}_{lesson.get('classroom', '')}"
                    if key not in seen:
                        seen.add(key)
                        unique_lessons.append(lesson)
                    else:
                        print(f"🧹 УДАЛЕН ДУБЛИКАТ: {key} в {day_key} {slot_key}")
                schedule[day_key][slot_key] = unique_lessons
    return schedule


def response_to_dict(result: TimetableResponse, request: TimetableRequest) -> Dict[str, Any]:
    return {
        "status": result.status,
        "solution_time": result.solution_time,
        "days": request.days,
        "slots": request.slots,
        "lessonTimes": request.lesson_times,
        "teachers": [{
            "teacher": t.teacher,
            "total_lessons": t.total_lessons,
            "schedule": {
                day: {
                    str(slot): [
                        {
                            "subject": a.subject,
                            "grade": a.grade,
                            "teacher": a.teacher,
                            "classroom": a.classroom,
                            "topic": "",
                            "week_index": a.week_index
                        }
                        for a in assignments
                    ]
                    for slot, assignments in slots.items()
                }
                for day, slots in t.schedule.items()
            }
        } for t in result.teachers],
        "validation_errors": result.validation_errors
    }


# ============================================================
# ПРИЛОЖЕНИЕ
# ============================================================
app = FastAPI(title="Генератор школьного расписания")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

app.mount("/static", StaticFiles(directory="static"), name="static")


@app.on_event("startup")
def startup_event():
    init_db()
    print("✅ База данных инициализирована")


@app.get("/", response_class=HTMLResponse)
async def read_root():
    index_path = os.path.join("static", "index.html")
    if os.path.exists(index_path):
        with open(index_path, "r", encoding="utf-8") as f:
            return f.read()
    return HTMLResponse(content="<h1>Frontend not found</h1>", status_code=404)


@app.get("/health")
def health_check():
    return {"status": "ok"}


# ============================================================
# АВТОРИЗАЦИЯ
# ============================================================
@app.post("/register")
def register(user: UserRegister, db: Session = Depends(get_db)):
    db_user = db.query(User).filter(User.email == user.email).first()
    if db_user:
        raise HTTPException(status_code=400, detail="Email уже зарегистрирован")
    new_user = User(
        email=user.email,
        hashed_password=get_password_hash(user.password),
        full_name=user.full_name,
        school_name=user.school_name
    )
    db.add(new_user)
    db.commit()
    db.refresh(new_user)
    access_token = create_access_token(
        data={"sub": new_user.id},
        expires_delta=timedelta(minutes=10080)
    )
    return {
        "message": "Регистрация успешна",
        "user_id": new_user.id,
        "access_token": access_token,
        "token_type": "bearer",
        "full_name": new_user.full_name
    }


@app.post("/login")
def login(user: UserLogin, db: Session = Depends(get_db)):
    db_user = db.query(User).filter(User.email == user.email).first()
    if not db_user or not verify_password(user.password, db_user.hashed_password):
        raise HTTPException(status_code=401, detail="Неверный email или пароль")
    access_token = create_access_token(
        data={"sub": db_user.id},
        expires_delta=timedelta(minutes=10080)
    )
    return {
        "message": "Вход выполнен успешно",
        "user_id": db_user.id,
        "full_name": db_user.full_name,
        "school_name": db_user.school_name,
        "access_token": access_token,
        "token_type": "bearer"
    }


# ============================================================
# AI ПАРСЕР
# ============================================================
@app.post("/api/ai-parse")
def ai_parse_text(request: AIInputRequest, current_user: User = Depends(get_current_user)):
    text = request.text.lower().strip()
    result = {"teacher_name": "", "lessons": [], "availability": {}, "max_days": 5}

    name_match = re.search(
        r'^(?:учитель\s+)?([а-яё]+(?:\s+[а-яё\.]+)*?)\s+(?:может|ведет|преподает|работает|проводит|занятия)',
        text
    )
    if name_match:
        raw_name = name_match.group(1).strip()
        stop_words = ['может', 'ведет', 'преподает', 'работает', 'проводит', 'занятия', 'урок', 'уроки', 'занятие']
        clean_name_parts = [w for w in raw_name.split() if w.lower() not in stop_words]
        result["teacher_name"] = " ".join(clean_name_parts).title()

    lesson_patterns = [
        r'(?:ведет|вести|проводит|проводить)\s+(?:урок\s+)?([а-яё\s]+?)\s+(?:в\s+)?(\d+(?:-\d+)?)\s*(?:класс|кл)',
        r'(?:преподает|преподавать)\s+([а-яё\s]+?)\s+(?:в\s+)?(\d+(?:-\d+)?)\s*(?:класс|кл)'
    ]
    for pattern in lesson_patterns:
        matches = re.findall(pattern, text)
        for match in matches:
            result["lessons"].append({"subject": match[0].strip().title(), "grade": match[1].strip(), "hours": 2})

    LESSON_TIMES = [
        ("08:30", "09:15"), ("09:25", "10:10"), ("10:20", "11:05"),
        ("11:25", "12:10"), ("12:20", "13:05"), ("13:15", "14:00"),
        ("14:10", "14:55"), ("15:05", "15:50")
    ]

    def get_slots_by_time(start_time_str=None, end_time_str=None):
        slots = []
        start_min = 0
        end_min = 24 * 60
        if start_time_str:
            clean = start_time_str.replace('.', ':').strip()
            parts = clean.split(':')
            if len(parts) == 2:
                try:
                    start_min = int(parts[0]) * 60 + int(parts[1])
                except:
                    pass
        if end_time_str:
            clean = end_time_str.replace('.', ':').replace('часов', '').replace('ч', '').strip()
            if ':' not in clean:
                clean += ":00"
            parts = clean.split(':')
            if len(parts) == 2:
                try:
                    end_min = int(parts[0]) * 60 + int(parts[1])
                except:
                    pass
        for i, (s, e) in enumerate(LESSON_TIMES):
            s_parts = s.split(':')
            e_parts = e.split(':')
            lesson_start = int(s_parts[0]) * 60 + int(s_parts[1])
            lesson_end = int(e_parts[0]) * 60 + int(e_parts[1])
            if lesson_start >= start_min and lesson_end <= end_min:
                slots.append(i + 1)
        return slots

    days_map = {
        "понедельник": "Понедельник", "вторник": "Вторник",
        "среда": "Среда", "четверг": "Четверг", "пятница": "Пятница",
        "суббота": "Суббота", "каждый день": "ALL"
    }
    found_days = []
    for day_ru, day_norm in days_map.items():
        if day_ru in text:
            if day_ru == "каждый день":
                found_days.extend(["Понедельник", "Вторник", "Среда", "Четверг", "Пятница"])
            else:
                found_days.append(day_norm)

    if found_days:
        slots = []
        range_match = re.search(r'с\s+(\d{1,2}[.:]\d{0,2})\s+до\s+(\d{1,2}[.:]\d{0,2}(?:\s*часов?)?)', text)
        if range_match:
            slots = get_slots_by_time(range_match.group(1), range_match.group(2))
        elif not slots:
            end_match = re.search(r'до\s+(\d{1,2}[.:]\d{0,2}(?:\s*часов?)?)', text)
            if end_match:
                slots = get_slots_by_time(end_time_str=end_match.group(1))
            elif not slots:
                start_match = re.search(r'после\s+(\d{1,2}[.:]\d{0,2}(?:\s*часов?)?)', text)
                if start_match:
                    slots = get_slots_by_time(start_time_str=start_match.group(1))
        target_days = set(found_days)
        if slots:
            for d in target_days:
                result["availability"][d] = slots

    if not result["availability"] and ("до" in text or "после" in text or "с" in text):
        slots = []
        range_match = re.search(r'с\s+(\d{1,2}[.:]\d{0,2})\s+до\s+(\d{1,2}[.:]\d{0,2})', text)
        if range_match:
            slots = get_slots_by_time(range_match.group(1), range_match.group(2))
        else:
            end_match = re.search(r'до\s+(\d{1,2}[.:]\d{0,2})', text)
            if end_match:
                slots = get_slots_by_time(end_time_str=end_match.group(1))
        if slots:
            for d in ["Понедельник", "Вторник", "Среда", "Четверг", "Пятница"]:
                result["availability"][d] = slots

    return result


# ============================================================
# ГЕНЕРАЦИЯ
# ============================================================
@app.post("/generate")
def generate_timetable(request: TimetableRequest, current_user: User = Depends(get_current_user),
                       db: Session = Depends(get_db)):
    try:
        print(f"🚀 НАЧАЛО ГЕНЕРАЦИИ для пользователя {current_user.id}")
        is_valid, teacher_stats, errors = validate_teachers(request)
        if not is_valid:
            return {
                "message": "Ошибки валидации",
                "status": "INVALID",
                "errors": errors,
                "teacher_stats": teacher_stats
            }

        print("️ Запускаю OR-Tools solver...")
        result: TimetableResponse = solve_timetable(request)
        print(f"✅ Solver завершил работу. Статус: {result.status}")

        if result.status not in ("OPTIMAL", "FEASIBLE"):
            return {
                "message": "Не удалось составить расписание",
                "status": result.status,
                "solution_time": result.solution_time,
                "validation_errors": result.validation_errors or ["Проверьте ограничения"]
            }

        data_dict = response_to_dict(result, request)
        saved = SavedTimetable(
            user_id=current_user.id,
            name=f"Расписание {datetime.utcnow().strftime('%d.%m.%Y %H:%M')}",
            data=data_dict,
        )
        db.add(saved)
        db.commit()
        db.refresh(saved)
        print(f"💾 Расписание сохранено с ID: {saved.id}")

        return {
            "message": "Расписание сгенерировано",
            "id": saved.id,
            "status": result.status,
            "solution_time": result.solution_time,
            "validation_errors": result.validation_errors
        }
    except Exception as e:
        print(f"❌ КРИТИЧЕСКАЯ ОШИБКА ГЕНЕРАЦИИ: {e}")
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=f"Ошибка генерации: {str(e)}")


# ============================================================
# РАСПИСАНИЯ — CRUD
# ============================================================
@app.get("/timetables")
def get_timetables(current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    try:
        db.expire_all()
        timetables = db.query(SavedTimetable).filter(
            SavedTimetable.user_id == current_user.id
        ).order_by(SavedTimetable.created_at.desc()).all()
        return [
            {
                "id": t.id,
                "name": t.name,
                "created_at": t.created_at.isoformat() if t.created_at else None,
                "data": ensure_topic_field(copy.deepcopy(t.data)) if t.data else {}
            }
            for t in timetables
        ]
    except Exception as e:
        print(f"❌ Ошибка получения расписаний: {e}")
        return []


@app.get("/timetables/{timetable_id}")
def get_timetable(timetable_id: int, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    try:
        db.expire_all()
        timetable = db.query(SavedTimetable).filter(
            SavedTimetable.id == timetable_id,
            SavedTimetable.user_id == current_user.id
        ).first()
        if not timetable:
            raise HTTPException(status_code=404, detail="Расписание не найдено")

        data = ensure_topic_field(copy.deepcopy(timetable.data)) if timetable.data else {}

        print(f"📥 GET /timetables/{timetable_id} - возвращаем данные с {len(data.get('teachers', []))} учителями")

        return {
            "id": timetable.id,
            "name": timetable.name,
            "created_at": timetable.created_at.isoformat() if timetable.created_at else None,
            "data": data
        }
    except HTTPException:
        raise
    except Exception as e:
        print(f"❌ Ошибка получения расписания {timetable_id}: {e}")
        raise HTTPException(status_code=500, detail=f"Ошибка: {str(e)}")


@app.put("/timetables/{timetable_id}")
def update_timetable(timetable_id: int, update_data: TimetableUpdate, current_user: User = Depends(get_current_user),
                     db: Session = Depends(get_db)):
    timetable = db.query(SavedTimetable).filter(
        SavedTimetable.id == timetable_id,
        SavedTimetable.user_id == current_user.id
    ).first()
    if not timetable:
        raise HTTPException(status_code=404, detail="Расписание не найдено")
    if update_data.name is not None:
        timetable.name = update_data.name
    if update_data.data is not None:
        timetable.data = copy.deepcopy(update_data.data)
    db.commit()
    db.refresh(timetable)
    return {"message": "Расписание обновлено", "id": timetable.id, "name": timetable.name}


@app.delete("/timetables/{timetable_id}")
def delete_timetable(timetable_id: int, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    timetable = db.query(SavedTimetable).filter(
        SavedTimetable.id == timetable_id,
        SavedTimetable.user_id == current_user.id
    ).first()
    if not timetable:
        raise HTTPException(status_code=404, detail="Расписание не найдено")
    db.delete(timetable)
    db.commit()
    return {"message": "Расписание удалено"}


# ============================================================
# ПЕРЕМЕЩЕНИЕ УРОКА (ИСПРАВЛЕННАЯ ВЕРСИЯ)
# ============================================================
@app.post("/timetables/{timetable_id}/move-lesson")
def move_lesson(
        timetable_id: int,
        move_data: dict,
        current_user: User = Depends(get_current_user),
        db: Session = Depends(get_db)
):
    """Перемещает урок в новое время с гарантированным сохранением в БД"""
    try:
        print(f"🔄 MOVE LESSON REQUEST: {move_data}")
        print(f"📋 Timetable ID: {timetable_id}, User ID: {current_user.id}")

        db.expire_all()

        timetable = db.query(SavedTimetable).filter(
            SavedTimetable.id == timetable_id,
            SavedTimetable.user_id == current_user.id
        ).first()

        if not timetable:
            print(f"⚠️ Расписание {timetable_id} не найдено для пользователя {current_user.id}")
            raise HTTPException(status_code=404, detail="Расписание не найдено")

        print(f"✅ Расписание загружено из БД")

        # Работаем с КОПИЕЙ данных: мутация ORM-объекта in-place не отслеживается SQLAlchemy,
        # поэтому изменения не сохранялись. Изменённую копию присваиваем в конце.
        data = copy.deepcopy(timetable.data)
        teachers = data.get('teachers', [])

        req_teacher = move_data['teacher'].strip().lower()
        req_subject = move_data['subject'].strip().lower()
        req_grade = str(move_data['grade']).strip()
        req_classroom = move_data.get('classroom', '').strip()
        req_week_index = move_data.get('week_index', None)
        if req_week_index is not None:
            try:
                req_week_index = int(req_week_index)
            except (TypeError, ValueError):
                req_week_index = None
        from_day = move_data['from_day']
        from_slot = str(move_data['from_slot'])  # ключи слотов в JSON — строки
        to_day = move_data['to_day']
        to_slot = str(move_data['to_slot'])

        print(
            f"📋 Параметры: teacher={req_teacher}, subject={req_subject}, grade={req_grade}, classroom={req_classroom}")
        print(f"📋 Перемещение: {from_day} {from_slot} -> {to_day} {to_slot}")

        target_teacher = None
        for t in teachers:
            if t['teacher'].strip().lower() == req_teacher:
                target_teacher = t
                break

        if not target_teacher:
            print(f"️ Учитель '{req_teacher}' не найден")
            raise HTTPException(status_code=404, detail="Учитель не найден в расписании")

        schedule = target_teacher.get('schedule', {})
        print(f"📅 Текущее расписание учителя: {list(schedule.keys())}")

        actual_from_day = None
        actual_to_day = None

        for day_key in schedule.keys():
            day_lower = day_key.strip().lower()
            if day_lower == from_day.lower():
                actual_from_day = day_key
            if day_lower == to_day.lower():
                actual_to_day = day_key

        if actual_to_day is None:
            actual_to_day = to_day
            print(f"ℹ️ Целевой день '{to_day}' не найден, будет создан")

        if actual_from_day is None:
            print(f"❌ Исходный день '{from_day}' не найден")
            raise HTTPException(status_code=404, detail=f"День {from_day} не найден")

        print(f"📅 actual_from_day={actual_from_day}, actual_to_day={actual_to_day}")

        # ✅ ШАГ 1: Удаляем урок со старого места
        day_schedule = schedule.get(actual_from_day, {})
        lessons_list = day_schedule.get(from_slot, [])

        print(f"🔍 Ищем урок для удаления в {actual_from_day} {from_slot}")
        print(f"📚 Уроков в слоте: {len(lessons_list)}")

        new_lessons_list = []
        removed = False

        for l in lessons_list:
            l_subject = str(l.get('subject', '')).strip().lower()
            l_grade = str(l.get('grade', '')).strip()
            l_classroom = str(l.get('classroom', '')).strip()

            print(f"  Проверяем: subject={l_subject}, grade={l_grade}, classroom={l_classroom}")

            # Удаляем только ПЕРВЫЙ найденный урок с совпадающими параметрами
            is_same = (l_subject == req_subject and l_grade == req_grade)

            if is_same and not removed:
                print(f"✅ НАЙДЕН И УДАЛЯЕМ: {l['subject']} {l['grade']} {l.get('classroom', '')}")
                removed = True
            else:
                new_lessons_list.append(l)

        if not removed:
            print(f"⚠️ Урок НЕ НАЙДЕН для удаления!")

        if new_lessons_list:
            day_schedule[from_slot] = new_lessons_list
            print(f"✅ Слот {from_slot} обновлен (осталось {len(new_lessons_list)} уроков)")
        else:
            if from_slot in day_schedule:
                del day_schedule[from_slot]
                print(f" Слот {from_slot} удален (пустой)")

        # ✅ ШАГ 2: Проверяем конфликты перед добавлением
        to_day_schedule = schedule.get(actual_to_day, {})

        if to_slot in to_day_schedule and to_day_schedule[to_slot]:
            print(f"❌ Конфликт: учитель уже занят в {actual_to_day} {to_slot}")
            raise HTTPException(
                status_code=409,
                detail=f"Конфликт: {move_data['teacher']} уже занят в {to_day} на {to_slot} уроке"
            )

        for t in teachers:
            if t['teacher'].strip().lower() == req_teacher:
                continue
            t_sched = t.get('schedule', {})

            for day_key, day_data in t_sched.items():
                if day_key.strip().lower() == to_day.lower():
                    for slot_key, lessons in day_data.items():
                        if str(slot_key).strip() == to_slot:
                            for lesson in lessons:
                                if str(lesson.get('grade', '')).strip() == req_grade:
                                    print(f"❌ Конфликт: класс {req_grade} уже имеет урок у {t['teacher']}")
                                    raise HTTPException(
                                        status_code=409,
                                        detail=f"Конфликт: Класс {req_grade} уже имеет урок в это время у учителя {t['teacher']}"
                                    )

        # ✅ ШАГ 3: Проверяем, нет ли уже такого урока в целевом слоте
        if actual_to_day not in schedule:
            schedule[actual_to_day] = {}
            print(f"✅ Создан день {actual_to_day}")

        if to_slot not in schedule[actual_to_day]:
            schedule[actual_to_day][to_slot] = []
            print(f"✅ Создан слот {to_slot} в {actual_to_day}")

        # Проверяем дубликаты перед добавлением
        existing_lessons = schedule[actual_to_day][to_slot]
        already_exists = False
        for existing in existing_lessons:
            e_subject = str(existing.get('subject', '')).strip().lower()
            e_grade = str(existing.get('grade', '')).strip()
            e_classroom = str(existing.get('classroom', '')).strip()

            if e_subject == req_subject and e_grade == req_grade and e_classroom == req_classroom:
                already_exists = True
                print(f"⚠️ Урок уже существует в целевом слоте: {req_subject} {req_grade} {req_classroom}")
                break

        if not already_exists:
            new_lesson = {
                'subject': move_data['subject'],
                'grade': move_data['grade'],
                'teacher': move_data['teacher'],
                'classroom': move_data.get('classroom', ''),
                'topic': '',
                'week_index': req_week_index
            }
            schedule[actual_to_day][to_slot].append(new_lesson)
            print(
                f"➕ ДОБАВЛЕН: {new_lesson['subject']} {new_lesson['grade']} {new_lesson.get('classroom', '')} в {actual_to_day} {to_slot}")
        else:
            print(f"⚠️ Урок НЕ ДОБАВЛЕН (уже существует)")

        # ✅ ШАГ 4: Удаляем дубликаты во всем расписании учителя
        schedule = remove_duplicates_from_schedule(schedule)

        # ✅ ШАГ 5: Сохраняем в БД
        timetable.data = data
        db.commit()

        print(f"✅ ДАННЫЕ СОХРАНЕНЫ В БД")

        db.expire_all()

        fresh_timetable = db.query(SavedTimetable).filter(
            SavedTimetable.id == timetable_id
        ).first()

        if not fresh_timetable:
            print(f"❌ КРИТИЧЕСКАЯ ОШИБКА: не удалось загрузить расписание после сохранения")
            raise HTTPException(status_code=500, detail="Ошибка сохранения")

        fresh_data = ensure_topic_field(copy.deepcopy(fresh_timetable.data))
        print(f"✅ Загружены актуальные данные из БД")
        print(f"📊 Количество учителей: {len(fresh_data.get('teachers', []))}")

        return {"message": "Урок успешно перемещен", "data": fresh_data}

    except HTTPException:
        raise
    except Exception as e:
        print(f"❌ CRITICAL ERROR in move_lesson: {e}")
        import traceback
        traceback.print_exc()
        db.rollback()
        raise HTTPException(status_code=500, detail=f"Внутренняя ошибка сервера: {str(e)}")


# ============================================================
# ОБНОВЛЕНИЕ ТЕМ УРОКОВ
# ============================================================
@app.post("/timetables/{timetable_id}/update-topics")
def update_lesson_topics(
        timetable_id: int,
        request: TopicsUpdateRequest,
        current_user: User = Depends(get_current_user),
        db: Session = Depends(get_db)
):
    """Массовое обновление тем уроков"""
    try:
        db.expire_all()

        timetable = db.query(SavedTimetable).filter(
            SavedTimetable.id == timetable_id,
            SavedTimetable.user_id == current_user.id
        ).first()

        if not timetable:
            raise HTTPException(status_code=404, detail="Расписание не найдено")

        # Работаем с КОПИЕЙ данных: мутация ORM-объекта in-place не отслеживается SQLAlchemy,
        # поэтому изменения не сохранялись. Изменённую копию присваиваем в конце.
        data = copy.deepcopy(timetable.data)
        teachers = data.get('teachers', [])
        updated_count = 0

        print(f"📝 Обновление тем: получено {len(request.updates)} обновлений")

        for update in request.updates:
            req_teacher = update.teacher.strip().lower()
            req_subject = update.subject.strip().lower()
            req_grade = str(update.grade).strip()
            target_day = update.day
            target_slot = str(update.slot)  # ключи слотов в JSON — строки

            print(
                f"  Обработка: teacher={req_teacher}, subject={req_subject}, grade={req_grade}, day={target_day}, slot={target_slot}")

            target_teacher = None
            for t in teachers:
                if t['teacher'].strip().lower() == req_teacher:
                    target_teacher = t
                    break

            if not target_teacher:
                print(f"  ⚠️ Учитель '{req_teacher}' не найден")
                continue

            schedule = target_teacher.get('schedule', {})
            if target_day not in schedule:
                print(f"  ⚠️ День '{target_day}' не найден в расписании")
                continue
            if target_slot not in schedule[target_day]:
                print(f"  ⚠️ Слот '{target_slot}' не найден в день '{target_day}'")
                continue

            lessons = schedule[target_day][target_slot]
            print(f"   Уроков в слоте: {len(lessons)}")

            for lesson in lessons:
                l_subject = str(lesson.get('subject', '')).strip().lower()
                l_grade = str(lesson.get('grade', '')).strip()

                print(f"    Проверяем урок: subject={l_subject}, grade={l_grade}")

                is_match = (l_subject == req_subject and l_grade == req_grade)

                if is_match:
                    print(f"    ✅ НАЙДЕН: обновляем тему на '{update.topic}'")
                    lesson['topic'] = update.topic
                    updated_count += 1
                    break
            else:
                print(f"    ⚠️ Урок не найден для обновления темы")

        print(f"✅ Обновлено тем: {updated_count}")

        timetable.data = data
        db.commit()
        db.expire_all()

        fresh_timetable = db.query(SavedTimetable).filter(
            SavedTimetable.id == timetable_id
        ).first()

        fresh_data = ensure_topic_field(copy.deepcopy(fresh_timetable.data)) if fresh_timetable else data

        print(f"✅ Темы сохранены в БД")

        return {
            "message": "Темы обновлены",
            "updated_count": updated_count,
            "data": fresh_data
        }

    except HTTPException:
        raise
    except Exception as e:
        print(f"❌ Ошибка обновления тем: {e}")
        import traceback
        traceback.print_exc()
        db.rollback()
        raise HTTPException(status_code=500, detail=f"Внутренняя ошибка: {str(e)}")


# ============================================================
# ЭКСПОРТ ТЕМ В EXCEL
# ============================================================
@app.get("/timetables/{timetable_id}/export/topics-excel")
def export_topics_excel(
        timetable_id: int,
        grade: str = Query(..., description="Класс для экспорта"),
        current_user: User = Depends(get_current_user),
        db: Session = Depends(get_db)
):
    """Экспорт тем уроков для конкретного класса в Excel"""
    try:
        db.expire_all()

        timetable = db.query(SavedTimetable).filter(
            SavedTimetable.id == timetable_id,
            SavedTimetable.user_id == current_user.id
        ).first()

        if not timetable:
            raise HTTPException(status_code=404, detail="Расписание не найдено")

        data = ensure_topic_field(copy.deepcopy(timetable.data)) if timetable.data else {}
        days = data.get('days', ['Понедельник', 'Вторник', 'Среда', 'Четверг', 'Пятница'])
        slots = data.get('slots', [1, 2, 3, 4, 5, 6, 7])
        times = data.get('lessonTimes', [])

        class_schedule = {}
        for day in days:
            class_schedule[day] = {}
            for slot in slots:
                class_schedule[day][slot] = None

        for teacher in data.get('teachers', []):
            if not teacher.get('schedule'):
                continue
            for day, day_data in teacher['schedule'].items():
                for slot_str, lessons in day_data.items():
                    slot = int(slot_str)
                    if not isinstance(lessons, list):
                        continue
                    for lesson in lessons:
                        if str(lesson.get('grade')) == str(grade):
                            class_schedule[day][slot] = {
                                'subject': lesson.get('subject', ''),
                                'teacher': lesson.get('teacher', ''),
                                'classroom': lesson.get('classroom', ''),
                                'topic': lesson.get('topic', '')
                            }

        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = f"Темы {grade}"

        header_font = Font(bold=True, size=11, color="FFFFFF")
        day_fill = PatternFill(start_color="F0F4FF", end_color="F0F4FF", fill_type="solid")
        header_fill = PatternFill(start_color="667EEA", end_color="667EEA", fill_type="solid")
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

        ws.cell(row=1, column=1, value=f"Расписание тем уроков — {grade} класс").font = Font(bold=True, size=14)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(slots) + 1)

        row = 3
        ws.cell(row=row, column=1, value="День / Урок").font = header_font
        ws.cell(row=row, column=1).fill = header_fill
        ws.cell(row=row, column=1).alignment = center_align
        ws.cell(row=row, column=1).border = thin_border

        for col_idx, slot in enumerate(slots, start=2):
            time_str = times[slot - 1] if slot - 1 < len(times) else ''
            cell = ws.cell(row=row, column=col_idx, value=f"Урок {slot}\n{time_str}")
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border

        for day_idx, day in enumerate(days):
            row = 4 + day_idx
            ws.cell(row=row, column=1, value=day).font = Font(bold=True)
            ws.cell(row=row, column=1).fill = day_fill
            ws.cell(row=row, column=1).alignment = center_align
            ws.cell(row=row, column=1).border = thin_border

            for col_idx, slot in enumerate(slots, start=2):
                lesson = class_schedule[day][slot]
                cell = ws.cell(row=row, column=col_idx)
                cell.border = thin_border
                cell.alignment = Alignment(vertical='top', wrap_text=True)

                if lesson:
                    cell.value = (
                        f"Предмет: {lesson['subject']}\n"
                        f"Учитель: {lesson['teacher']}\n"
                        f"Кабинет: {lesson['classroom']}\n"
                        f"Тема: {lesson['topic']}"
                    )
                else:
                    cell.value = "—"

        ws.column_dimensions['A'].width = 15
        for col_idx in range(2, len(slots) + 2):
            col_letter = chr(64 + col_idx) if col_idx <= 26 else 'A'
            ws.column_dimensions[col_letter].width = 30

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        filename = f"temy_{grade}_{timetable_id}.xlsx"
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )

    except HTTPException:
        raise
    except Exception as e:
        print(f"❌ Ошибка экспорта тем: {e}")
        raise HTTPException(status_code=500, detail=f"Ошибка: {str(e)}")


# ============================================================
# ЭКСПОРТ EXCEL
# ============================================================
@app.get("/timetables/{timetable_id}/export/excel")
def export_excel(
        timetable_id: int,
        format: str = Query("classes", pattern="^(teachers|classes)$"),
        current_user: User = Depends(get_current_user),
        db: Session = Depends(get_db)
):
    db.expire_all()

    timetable = db.query(SavedTimetable).filter(
        SavedTimetable.id == timetable_id,
        SavedTimetable.user_id == current_user.id
    ).first()

    if not timetable:
        raise HTTPException(status_code=404, detail="Расписание не найдено")

    data = ensure_topic_field(copy.deepcopy(timetable.data)) if timetable.data else {}
    if 'lessonTimes' not in data:
        data['lessonTimes'] = [
            "08:30-09:15", "09:25-10:10", "10:20-11:05",
            "11:25-12:10", "12:20-13:05", "13:15-14:00", "14:10-14:55"
        ]

    wb = generate_class_template_excel(data)
    filename = f"raspisanie_{format}_{timetable_id}.xlsx"
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


# ============================================================
# ЭКСПОРТ PDF
# ============================================================
@app.get("/timetables/{timetable_id}/export/pdf")
def export_pdf(
        timetable_id: int,
        format: str = Query("teachers", pattern="^(teachers|classes)$"),
        current_user: User = Depends(get_current_user),
        db: Session = Depends(get_db)
):
    db.expire_all()

    timetable = db.query(SavedTimetable).filter(
        SavedTimetable.id == timetable_id,
        SavedTimetable.user_id == current_user.id
    ).first()

    if not timetable:
        raise HTTPException(status_code=404, detail="Расписание не найдено")

    data = ensure_topic_field(copy.deepcopy(timetable.data)) if timetable.data else {}
    days = data.get("days", ["Понедельник", "Вторник", "Среда", "Четверг", "Пятница"])
    slots = data.get("slots", [1, 2, 3, 4, 5, 6, 7])

    if format == "teachers":
        pdf_bytes = generate_teachers_pdf(data, days, slots)
        filename = f"timetable_teachers_{timetable_id}.pdf"
    else:
        pdf_bytes = generate_classes_pdf(data, days, slots)
        filename = f"timetable_classes_{timetable_id}.pdf"

    return StreamingResponse(
        io.BytesIO(pdf_bytes),
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


# ============================================================
# ЗАМЕНЫ (общий журнал для всех пользователей)
# ============================================================
class SubstitutionCreate(BaseModel):
    date: Optional[str] = None
    slot: int = 1
    grade: Optional[str] = None
    absent: str
    substitute: Optional[str] = None


class SubstitutionUpdate(BaseModel):
    substitute: Optional[str] = None


def _substitution_to_dict(s: Substitution) -> Dict[str, Any]:
    return {
        "id": s.id,
        "date": s.date,
        "slot": s.slot,
        "grade": s.grade,
        "absent": s.absent,
        "substitute": s.substitute or "",
        "created_by_name": s.created_by_name,
        "created_at": s.created_at.isoformat() if s.created_at else None,
    }


@app.get("/api/substitutions")
def list_substitutions(current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    items = db.query(Substitution).order_by(Substitution.date.desc(), Substitution.slot.asc()).all()
    return [_substitution_to_dict(s) for s in items]


@app.post("/api/substitutions")
def create_substitution(payload: SubstitutionCreate, current_user: User = Depends(get_current_user),
                        db: Session = Depends(get_db)):
    if payload.substitute and payload.substitute.strip() == payload.absent.strip():
        raise HTTPException(status_code=400, detail="Замена не может совпадать с отсутствующим учителем")
    s = Substitution(
        date=(payload.date or None),
        slot=payload.slot or 1,
        grade=payload.grade,
        absent=payload.absent,
        substitute=payload.substitute or "",
        created_by_name=current_user.full_name,
    )
    db.add(s)
    db.commit()
    db.refresh(s)
    return _substitution_to_dict(s)


@app.patch("/api/substitutions/{sub_id}")
def update_substitution(sub_id: int, payload: SubstitutionUpdate,
                        current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    s = db.query(Substitution).filter(Substitution.id == sub_id).first()
    if not s:
        raise HTTPException(status_code=404, detail="Замена не найдена")
    sub_value = (payload.substitute or "").strip()
    if sub_value and sub_value == (s.absent or "").strip():
        raise HTTPException(status_code=400, detail="Замена не может совпадать с отсутствующим учителем")
    s.substitute = sub_value
    db.commit()
    db.refresh(s)
    return _substitution_to_dict(s)


@app.delete("/api/substitutions/{sub_id}")
def delete_substitution(sub_id: int, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    s = db.query(Substitution).filter(Substitution.id == sub_id).first()
    if not s:
        raise HTTPException(status_code=404, detail="Замена не найдена")
    db.delete(s)
    db.commit()
    return {"message": "Замена удалена"}


# ============================================================
# ИМПОРТ СПРАВОЧНИКОВ ИЗ EXCEL
# ============================================================
_HEADER_WORDS = (
    "класс", "class", "учит", "teacher", "фио", "предмет", "subject",
    "кабинет", "room", "номер", "название", "категор", "вместим",
)


def _norm_cell(value) -> str:
    return str(value).strip() if value is not None else ""


def _looks_like_header(row) -> bool:
    for cell in row:
        s = _norm_cell(cell).lower()
        if s and any(w in s for w in _HEADER_WORDS):
            return True
    return False


def _sheet_target(title: str):
    t = (title or "").lower()
    if "класс" in t or "class" in t:
        return "classes"
    if "учит" in t or "teacher" in t:
        return "teachers"
    if "предмет" in t or "subject" in t:
        return "subjects"
    if "кабинет" in t or "room" in t:
        return "classrooms"
    return None


def _parse_reference_workbook(wb) -> Dict[str, Any]:
    """Разбирает книгу Excel в списки классов/учителей/предметов/кабинетов.

    Поддерживаются два варианта:
    1) отдельные листы «Классы», «Учителя», «Предметы», «Кабинеты»;
    2) один лист с колонками, где заголовки содержат «Класс», «Учитель», «Предмет», «Кабинет».
    """
    result = {"classes": [], "teachers": [], "subjects": [], "classrooms": []}

    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        target = _sheet_target(ws.title)
        if not target:
            continue
        start = 1 if _looks_like_header(rows[0]) else 0
        for r in rows[start:]:
            name = _norm_cell(r[0]) if len(r) > 0 else ""
            if not name:
                continue
            if target == "subjects":
                cat = _norm_cell(r[1]) if len(r) > 1 else ""
                result["subjects"].append({"name": name, "category": cat or "Основной"})
            elif target == "classrooms":
                cap_raw = r[1] if len(r) > 1 else None
                try:
                    cap = int(cap_raw) if cap_raw not in (None, "") else None
                except Exception:
                    cap = None
                result["classrooms"].append({"number": name, "capacity": cap})
            else:
                result[target].append(name)

    if not any(result.values()) and wb.worksheets:
        ws = wb.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        if rows:
            header = [_norm_cell(h).lower() for h in rows[0]]

            def col(words):
                for i, h in enumerate(header):
                    if any(w in h for w in words):
                        return i
                return None

            ci = col(["класс", "class"])
            ti = col(["учит", "teacher", "фио"])
            si = col(["предмет", "subject"])
            ri = col(["кабинет", "room"])

            for r in rows[1:]:
                if ci is not None and ci < len(r) and _norm_cell(r[ci]):
                    result["classes"].append(_norm_cell(r[ci]))
                if ti is not None and ti < len(r) and _norm_cell(r[ti]):
                    result["teachers"].append(_norm_cell(r[ti]))
                if si is not None and si < len(r) and _norm_cell(r[si]):
                    result["subjects"].append({"name": _norm_cell(r[si]), "category": "Основной"})
                if ri is not None and ri < len(r) and _norm_cell(r[ri]):
                    result["classrooms"].append({"number": _norm_cell(r[ri]), "capacity": None})

    result["classes"] = list(dict.fromkeys(result["classes"]))
    result["teachers"] = list(dict.fromkeys(result["teachers"]))

    seen = set()
    subs = []
    for s in result["subjects"]:
        if s["name"] not in seen:
            seen.add(s["name"])
            subs.append(s)
    result["subjects"] = subs

    seen = set()
    rooms = []
    for c in result["classrooms"]:
        if c["number"] not in seen:
            seen.add(c["number"])
            rooms.append(c)
    result["classrooms"] = rooms

    return result


@app.post("/api/import-excel")
async def import_excel(file: UploadFile = File(...), current_user: User = Depends(get_current_user)):
    """Импорт справочников из Excel: возвращает распознанные списки."""
    try:
        content = await file.read()
        if not content:
            raise HTTPException(status_code=400, detail="Пустой файл")
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
        try:
            data = _parse_reference_workbook(wb)
        finally:
            wb.close()
        return data
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Не удалось прочитать Excel: {e}")


@app.get("/api/import-template")
def import_template(current_user: User = Depends(get_current_user)):
    """Скачать пустой шаблон Excel для импорта справочников."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")

    def make_sheet(ws, title, headers, rows):
        ws.title = title
        ws.append(headers)
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=c)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        for r in rows:
            ws.append(r)
        for col in ws.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None), default=10)
            ws.column_dimensions[col[0].column_letter].width = max(14, width + 4)

    ws1 = wb.active
    make_sheet(ws1, "Классы", ["Класс"], [["5А"], ["6Б"]])

    ws2 = wb.create_sheet()
    make_sheet(ws2, "Учителя", ["Учитель"], [["Иванов И.И."], ["Петров С.П."]])

    ws3 = wb.create_sheet()
    make_sheet(ws3, "Предметы", ["Предмет", "Категория"],
               [["Математика", "Основной"], ["Физика", "Факультатив"]])

    ws4 = wb.create_sheet()
    make_sheet(ws4, "Кабинеты", ["Кабинет", "Вместимость"], [["101", 30], ["205", 25]])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=shablon_importa.xlsx"},
    )