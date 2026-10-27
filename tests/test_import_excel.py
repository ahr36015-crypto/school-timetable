"""Тесты импорта справочников из Excel."""
import io
import uuid

from openpyxl import Workbook


def _register(client):
    email = f"imp_{uuid.uuid4().hex}@test.ru"
    r = client.post("/register", json={
        "email": email, "password": "secret123", "full_name": "Imp", "school_name": "S",
    })
    assert r.status_code == 200, r.text
    return {"Authorization": "Bearer " + r.json()["access_token"]}


def _xlsx_bytes(wb):
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_import_excel_parses_separate_sheets(client):
    wb = Workbook()
    ws = wb.active
    ws.title = "Классы"
    ws.append(["Класс"])
    ws.append(["5А"])
    ws.append(["6Б"])

    wt = wb.create_sheet("Учителя")
    wt.append(["Учитель"])
    wt.append(["Иванов И.И."])
    wt.append(["Петров С.П."])

    wsp = wb.create_sheet("Предметы")
    wsp.append(["Предмет", "Категория"])
    wsp.append(["Математика", "Основной"])
    wsp.append(["Физика", "Факультатив"])

    wr = wb.create_sheet("Кабинеты")
    wr.append(["Кабинет", "Вместимость"])
    wr.append(["101", 30])
    wr.append(["205", 25])

    headers = _register(client)
    r = client.post(
        "/api/import-excel", headers=headers,
        files={"file": ("refs.xlsx", _xlsx_bytes(wb),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert r.status_code == 200, r.text
    d = r.json()

    assert d["classes"] == ["5А", "6Б"]
    assert d["teachers"] == ["Иванов И.И.", "Петров С.П."]
    assert {"name": "Математика", "category": "Основной"} in d["subjects"]
    assert {"name": "Физика", "category": "Факультатив"} in d["subjects"]
    assert any(c["number"] == "101" and c["capacity"] == 30 for c in d["classrooms"])
    assert any(c["number"] == "205" for c in d["classrooms"])


def test_import_excel_single_sheet_with_columns(client):
    wb = Workbook()
    ws = wb.active
    ws.title = "Справочники"
    ws.append(["Класс", "Учитель", "Предмет", "Кабинет"])
    ws.append(["7А", "Сидоров С.С.", "Химия", "301"])

    headers = _register(client)
    r = client.post(
        "/api/import-excel", headers=headers,
        files={"file": ("refs.xlsx", _xlsx_bytes(wb),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert r.status_code == 200, r.text
    d = r.json()

    assert "7А" in d["classes"]
    assert "Сидоров С.С." in d["teachers"]
    assert any(s["name"] == "Химия" for s in d["subjects"])
    assert any(c["number"] == "301" for c in d["classrooms"])


def test_import_excel_requires_auth(client):
    r = client.post("/api/import-excel",
                    files={"file": ("x.xlsx", b"x", "application/octet-stream")})
    assert r.status_code in (401, 403)
